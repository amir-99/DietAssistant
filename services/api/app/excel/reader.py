"""Reads diet plan data from the active workbook."""
from __future__ import annotations
import re
from typing import List, Optional
import openpyxl
from app.models.diet import DietOption, ParsedItem, PlanSectionSummary

SECTION_SHEETS = ["Breakfast", "Lunch", "Dinner", "Snack 1", "Snack 2", "Snack 3"]


def _col_index(ws, name: str) -> Optional[int]:
    for idx, cell in enumerate(ws[1], start=1):
        if cell.value and str(cell.value).strip() == name:
            return idx
    return None


def read_plan_sections(path: str) -> List[PlanSectionSummary]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    summaries: List[PlanSectionSummary] = []

    if "Index" in wb.sheetnames:
        ws = wb["Index"]
        sec_col = _col_index(ws, "Section")
        cnt_col = _col_index(ws, "Options Extracted")
        page_col = _col_index(ws, "PDF Page(s)")
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            section = str(row[(sec_col or 1) - 1]).strip() if sec_col else ""
            count_raw = row[(cnt_col or 2) - 1] if cnt_col else 0
            pages = str(row[(page_col or 3) - 1]).strip() if (page_col and row[(page_col - 1)]) else ""
            count = int(count_raw) if count_raw else 0
            if section:
                summaries.append(PlanSectionSummary(section=section, option_count=count, pdf_pages=pages))
    else:
        # Fall back to reading each section sheet
        for sheet in SECTION_SHEETS:
            if sheet in wb.sheetnames:
                ws = wb[sheet]
                count = sum(1 for row in ws.iter_rows(min_row=2, values_only=True) if row and row[0])
                summaries.append(PlanSectionSummary(section=sheet, option_count=count))

    wb.close()
    return summaries


def read_all_options(path: str) -> List[DietOption]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    options: List[DietOption] = []

    if "All Sections" in wb.sheetnames:
        ws = wb["All Sections"]
        sec_col = _col_index(ws, "Section")
        opt_col = _col_index(ws, "Option No.")
        txt_col = _col_index(ws, "Diet option (exact PDF text)")
        page_col = _col_index(ws, "PDF Page")

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            section = str(row[(sec_col or 1) - 1]).strip() if sec_col and row[(sec_col - 1)] else ""
            opt_no_raw = row[(opt_col or 2) - 1] if opt_col else None
            text = str(row[(txt_col or 3) - 1]).strip() if txt_col and row[(txt_col - 1)] else ""
            page = str(row[(page_col or 4) - 1]).strip() if page_col and row[(page_col - 1)] else ""
            if not section or opt_no_raw is None or not text:
                continue
            try:
                opt_no = int(opt_no_raw)
            except (ValueError, TypeError):
                continue
            options.append(DietOption(section=section, option_no=opt_no, option_text=text, pdf_page=page or None))
    else:
        # Fall back to per-section sheets
        for sheet in SECTION_SHEETS:
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            opt_col = _col_index(ws, "Option No.")
            txt_col = _col_index(ws, "Diet option (exact PDF text)")
            page_col = _col_index(ws, "PDF Page")
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row:
                    continue
                opt_no_raw = row[(opt_col or 1) - 1] if opt_col else None
                text = str(row[(txt_col or 2) - 1]).strip() if txt_col and row[(txt_col - 1)] else ""
                page = str(row[(page_col or 3) - 1]).strip() if page_col and row[(page_col - 1)] else ""
                if opt_no_raw is None or not text:
                    continue
                try:
                    opt_no = int(opt_no_raw)
                except (ValueError, TypeError):
                    continue
                options.append(DietOption(section=sheet, option_no=opt_no, option_text=text, pdf_page=page or None))

    wb.close()
    return options


def read_consumption_log(path: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    if "Consumption_Log" not in wb.sheetnames:
        wb.close()
        return rows

    ws = wb["Consumption_Log"]
    headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        rows.append(dict(zip(headers, row)))
    wb.close()
    return rows


def read_parsed_items(path: str) -> List[ParsedItem]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    items: List[ParsedItem] = []
    if "Parsed_Option_Items" not in wb.sheetnames:
        wb.close()
        return items

    ws = wb["Parsed_Option_Items"]
    headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        if not row or not row[0]:
            continue
        d = dict(zip(headers, row))
        try:
            items.append(ParsedItem(
                id=i,
                section=str(d.get("section", "")),
                option_no=int(d.get("option_no", 0)),
                source_option_text=str(d.get("source_option_text", "")),
                parsed_item_order=int(d.get("parsed_item_order", i)),
                item_name_guess=str(d.get("item_name_guess", "")),
                quantity_text=str(d.get("quantity_text", "")),
                quantity_numeric=_safe_float(d.get("quantity_numeric")),
                unit_guess=str(d.get("unit_guess", "")),
                parser_confidence=float(d.get("parser_confidence", 0)),
                review_status=str(d.get("review_status", "auto")),
            ))
        except Exception:
            continue
    wb.close()
    return items


def _safe_float(v) -> Optional[float]:
    try:
        return float(v)
    except (TypeError, ValueError):
        return None
