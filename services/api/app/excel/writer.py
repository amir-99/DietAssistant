"""Writes tracking data into the active workbook with file locking."""
from __future__ import annotations
import uuid
import shutil
from datetime import datetime
from pathlib import Path
from typing import List, Optional

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from filelock import FileLock

from app.models.tracking import ConsumptionEventCreate, EventPatch
from app.models.diet import ParsedItem

CONSUMPTION_HEADERS = [
    "event_id", "timestamp_local", "date_local", "raw_user_message",
    "log_type", "section", "option_no", "source_option_text",
    "item_name", "consumed_qty_text", "consumed_qty_numeric",
    "consumed_unit", "portion_fraction", "confidence", "status",
    "estimated_calories", "notes",
]

PARSED_HEADERS = [
    "section", "option_no", "source_option_text", "parsed_item_order",
    "item_name_guess", "quantity_text", "quantity_numeric",
    "unit_guess", "parser_confidence", "review_status",
]

DAILY_HEADERS = [
    "date_local", "status_type", "section", "option_no", "item_name",
    "planned_qty_text", "consumed_qty_text", "consumed_qty_numeric",
    "remaining_qty_numeric", "completion_pct", "confidence_note",
]

PINK_FILL = PatternFill("solid", fgColor="FFB6C1")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="E91E8C")


def _lock_path(workbook_path: str) -> str:
    return workbook_path + ".lock"


def _style_header_row(ws) -> None:
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def ensure_tracking_sheets(workbook_path: str) -> None:
    lock = FileLock(_lock_path(workbook_path), timeout=30)
    with lock:
        wb = openpyxl.load_workbook(workbook_path)
        changed = False
        if "Consumption_Log" not in wb.sheetnames:
            ws = wb.create_sheet("Consumption_Log")
            ws.append(CONSUMPTION_HEADERS)
            _style_header_row(ws)
            changed = True
        if "Parsed_Option_Items" not in wb.sheetnames:
            ws = wb.create_sheet("Parsed_Option_Items")
            ws.append(PARSED_HEADERS)
            _style_header_row(ws)
            changed = True
        if "Daily_Status" not in wb.sheetnames:
            ws = wb.create_sheet("Daily_Status")
            ws.append(DAILY_HEADERS)
            _style_header_row(ws)
            changed = True
        if changed:
            wb.save(workbook_path)
        wb.close()


def write_consumption_event(
    workbook_path: str,
    entry: ConsumptionEventCreate,
    raw_message: str,
    date_local: str,
    tz: str = "Asia/Tehran",
) -> str:
    import pytz
    from datetime import timezone
    tz_obj = pytz.timezone(tz)
    now = datetime.now(tz_obj)
    event_id = str(uuid.uuid4())[:8]
    timestamp = now.strftime("%Y-%m-%d %H:%M:%S")

    row = [
        event_id, timestamp, date_local, raw_message,
        entry.log_type,
        entry.section or "",
        entry.option_no or "",
        entry.source_option_text or "",
        entry.item_name or "",
        entry.consumed_qty_text or "",
        entry.consumed_qty_numeric if entry.consumed_qty_numeric is not None else "",
        entry.consumed_unit or "",
        entry.portion_fraction if entry.portion_fraction is not None else "",
        entry.confidence if entry.confidence is not None else "",
        entry.status,
        entry.estimated_calories if entry.estimated_calories is not None else "",
        entry.notes or "",
    ]

    lock = FileLock(_lock_path(workbook_path), timeout=30)
    with lock:
        wb = openpyxl.load_workbook(workbook_path)
        ensure_log_sheet(wb)
        ws = wb["Consumption_Log"]
        ws.append(row)
        wb.save(workbook_path)
        wb.close()

    return event_id


def patch_consumption_event(workbook_path: str, event_id: str, patch: EventPatch) -> bool:
    lock = FileLock(_lock_path(workbook_path), timeout=30)
    with lock:
        wb = openpyxl.load_workbook(workbook_path)
        if "Consumption_Log" not in wb.sheetnames:
            wb.close()
            return False
        ws = wb["Consumption_Log"]
        headers = [str(c.value) for c in ws[1]]
        id_col = headers.index("event_id") + 1 if "event_id" in headers else None
        if not id_col:
            wb.close()
            return False

        found = False
        for row in ws.iter_rows(min_row=2):
            if row[id_col - 1].value == event_id:
                patch_dict = patch.model_dump(exclude_none=True)
                for field, value in patch_dict.items():
                    if field in headers:
                        col = headers.index(field) + 1
                        row[col - 1].value = value
                if patch.status is None:
                    status_col = headers.index("status") + 1 if "status" in headers else None
                    if status_col:
                        row[status_col - 1].value = "corrected"
                found = True
                break

        if found:
            wb.save(workbook_path)
        wb.close()
        return found


def delete_consumption_event(workbook_path: str, event_id: str) -> bool:
    lock = FileLock(_lock_path(workbook_path), timeout=30)
    with lock:
        wb = openpyxl.load_workbook(workbook_path)
        if "Consumption_Log" not in wb.sheetnames:
            wb.close()
            return False
        ws = wb["Consumption_Log"]
        headers = [str(c.value) for c in ws[1]]
        id_col = headers.index("event_id") + 1 if "event_id" in headers else None
        status_col = headers.index("status") + 1 if "status" in headers else None
        if not id_col:
            wb.close()
            return False

        found = False
        for row in ws.iter_rows(min_row=2):
            if row[id_col - 1].value == event_id:
                if status_col:
                    row[status_col - 1].value = "deleted"
                found = True
                break

        if found:
            wb.save(workbook_path)
        wb.close()
        return found


def write_parsed_items(workbook_path: str, items: List[ParsedItem]) -> None:
    lock = FileLock(_lock_path(workbook_path), timeout=30)
    with lock:
        wb = openpyxl.load_workbook(workbook_path)
        if "Parsed_Option_Items" in wb.sheetnames:
            del wb["Parsed_Option_Items"]
        ws = wb.create_sheet("Parsed_Option_Items")
        ws.append(PARSED_HEADERS)
        _style_header_row(ws)
        for item in items:
            ws.append([
                item.section, item.option_no, item.source_option_text,
                item.parsed_item_order, item.item_name_guess,
                item.quantity_text, item.quantity_numeric, item.unit_guess,
                item.parser_confidence, item.review_status,
            ])
        wb.save(workbook_path)
        wb.close()


def backup_workbook(workbook_path: str, backup_dir: str) -> str:
    from datetime import datetime
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    dest = Path(backup_dir) / f"diet_plan_{ts}.xlsx"
    Path(backup_dir).mkdir(parents=True, exist_ok=True)
    shutil.copy2(workbook_path, dest)
    return str(dest)


def ensure_log_sheet(wb: Workbook) -> None:
    if "Consumption_Log" not in wb.sheetnames:
        ws = wb.create_sheet("Consumption_Log")
        ws.append(CONSUMPTION_HEADERS)
        _style_header_row(ws)
