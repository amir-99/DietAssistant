"""Generates a downloadable Excel template."""
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

PINK_HEADER = PatternFill("solid", fgColor="E91E8C")
PINK_LIGHT = PatternFill("solid", fgColor="FFB6C1")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _write_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = PINK_HEADER
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def generate_template() -> bytes:
    wb = openpyxl.Workbook()

    # All Sections
    ws = wb.active
    ws.title = "All Sections"
    _write_header(ws, ["Section", "Option No.", "Diet option (exact PDF text)", "PDF Page"])
    sample_row = ["Breakfast", 1, "Example: ۲ عدد نان تست + ۱ لیوان شیر + ۲ عدد خرما", "1"]
    ws.append(sample_row)
    ws.column_dimensions["C"].width = 60

    # Section sheets
    for section in ["Breakfast", "Lunch", "Dinner", "Snack 1", "Snack 2", "Snack 3"]:
        ws2 = wb.create_sheet(section)
        _write_header(ws2, ["Option No.", "Diet option (exact PDF text)", "PDF Page"])
        ws2.column_dimensions["B"].width = 60

    # Index
    ws_idx = wb.create_sheet("Index")
    _write_header(ws_idx, ["Section", "Options Extracted", "PDF Page(s)"])
    for section in ["Breakfast", "Lunch", "Dinner", "Snack 1", "Snack 2", "Snack 3"]:
        ws_idx.append([section, 0, ""])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
